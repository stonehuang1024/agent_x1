"""
Unit Tests for Agent X1 Tools Library

Tests all tool categories without requiring external API keys or network access.
Uses temporary files/directories for file/PDF/PPT/data tools.
"""

import sys
import os
import json
import tempfile
import unittest
from pathlib import Path

# Ensure project root is on path
sys.path.insert(0, str(Path(__file__).parent.parent.parent))


class TestToolRegistry(unittest.TestCase):
    """Tests for CategorizedToolRegistry and __init__ wiring."""

    def test_all_tools_import(self):
        from src.tools import ALL_TOOLS, TOOL_REGISTRY
        self.assertGreater(len(ALL_TOOLS), 30)
        self.assertEqual(len(ALL_TOOLS), len(TOOL_REGISTRY))

    def test_registry_categories(self):
        from src.tools import TOOL_REGISTRY
        catalog = TOOL_REGISTRY.get_catalog()
        expected = {"utility", "search", "stock", "economics", "file", "bash", "pdf", "ppt", "web", "data"}
        self.assertTrue(expected.issubset(set(catalog.keys())))

    def test_registry_search(self):
        from src.tools import TOOL_REGISTRY
        results = TOOL_REGISTRY.search("pdf")
        self.assertGreater(len(results), 0)
        for r in results:
            self.assertIn("pdf", (r["name"] + r["description"]).lower())

    def test_registry_search_file(self):
        from src.tools import TOOL_REGISTRY
        results = TOOL_REGISTRY.search("file")
        self.assertGreater(len(results), 0)

    def test_registry_get_tool(self):
        from src.tools import TOOL_REGISTRY
        tool = TOOL_REGISTRY.get("read_file")
        self.assertIsNotNone(tool)
        self.assertEqual(tool.name, "read_file")

    def test_registry_get_unknown_tool(self):
        from src.tools import TOOL_REGISTRY
        self.assertIsNone(TOOL_REGISTRY.get("nonexistent_tool_xyz"))

    def test_registry_category_lookup(self):
        from src.tools import TOOL_REGISTRY
        self.assertEqual(TOOL_REGISTRY.get_category("read_file"), "file")
        self.assertEqual(TOOL_REGISTRY.get_category("run_command"), "bash")
        self.assertEqual(TOOL_REGISTRY.get_category("read_pdf"), "pdf")
        self.assertEqual(TOOL_REGISTRY.get_category("create_presentation"), "ppt")
        self.assertEqual(TOOL_REGISTRY.get_category("fetch_url"), "web")
        self.assertEqual(TOOL_REGISTRY.get_category("read_csv"), "data")
        self.assertEqual(TOOL_REGISTRY.get_category("get_fred_series"), "economics")

    def test_all_tools_have_schemas(self):
        from src.tools import ALL_TOOLS
        for tool in ALL_TOOLS:
            schema = tool.get_schema()
            self.assertIn("type", schema)
            self.assertIn("function", schema)
            self.assertIn("name", schema["function"])
            self.assertIn("description", schema["function"])
            self.assertIn("parameters", schema["function"])

    def test_tool_names_are_unique(self):
        from src.tools import ALL_TOOLS
        names = [t.name for t in ALL_TOOLS]
        self.assertEqual(len(names), len(set(names)), "Duplicate tool names found")


class TestFileTools(unittest.TestCase):
    """Tests for file_tools.py."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_write_and_read_file(self):
        from src.tools.file_tools import write_file, read_file
        path = os.path.join(self.tmp, "test.txt")
        content = "Hello, Agent X1!\nLine 2"

        result = write_file(path, content)
        self.assertTrue(result["success"])
        self.assertEqual(result["lines"], 2)

        read_result = read_file(path)
        self.assertNotIn("error", read_result)
        self.assertEqual(read_result["content"], content)
        self.assertEqual(read_result["lines"], 2)

    def test_read_nonexistent_file(self):
        from src.tools.file_tools import read_file
        result = read_file("/nonexistent/path/file.txt")
        self.assertIn("error", result)

    def test_write_file_no_overwrite(self):
        from src.tools.file_tools import write_file
        path = os.path.join(self.tmp, "noover.txt")
        write_file(path, "original")
        result = write_file(path, "new", overwrite=False)
        self.assertIn("error", result)

    def test_append_file(self):
        from src.tools.file_tools import append_file, read_file
        path = os.path.join(self.tmp, "append.txt")
        append_file(path, "first\n")
        append_file(path, "second\n")
        r = read_file(path)
        self.assertIn("first", r["content"])
        self.assertIn("second", r["content"])

    def test_list_directory(self):
        from src.tools.file_tools import write_file, list_directory
        for i in range(3):
            write_file(os.path.join(self.tmp, f"f{i}.txt"), f"content{i}")
        result = list_directory(self.tmp)
        self.assertNotIn("error", result)
        self.assertEqual(result["count"], 3)

    def test_list_directory_pattern(self):
        from src.tools.file_tools import write_file, list_directory
        write_file(os.path.join(self.tmp, "a.txt"), "a")
        write_file(os.path.join(self.tmp, "b.py"), "b")
        result = list_directory(self.tmp, pattern="*.txt")
        self.assertEqual(result["count"], 1)
        self.assertEqual(result["entries"][0]["name"], "a.txt")

    def test_search_in_files(self):
        from src.tools.file_tools import write_file, search_in_files
        write_file(os.path.join(self.tmp, "a.txt"), "find me here\nnothing")
        write_file(os.path.join(self.tmp, "b.txt"), "not here")
        result = search_in_files(self.tmp, "find me", pattern="*.txt")
        self.assertEqual(result["total_matches"], 1)
        self.assertEqual(result["results"][0]["line"], 1)

    def test_move_file(self):
        from src.tools.file_tools import write_file, move_file, read_file
        src = os.path.join(self.tmp, "src.txt")
        dst = os.path.join(self.tmp, "dst.txt")
        write_file(src, "move me")
        result = move_file(src, dst)
        self.assertTrue(result["success"])
        self.assertFalse(os.path.exists(src))
        read_result = read_file(dst)
        self.assertEqual(read_result["content"], "move me")

    def test_copy_file(self):
        from src.tools.file_tools import write_file, copy_file
        src = os.path.join(self.tmp, "orig.txt")
        dst = os.path.join(self.tmp, "copy.txt")
        write_file(src, "copy me")
        result = copy_file(src, dst)
        self.assertTrue(result["success"])
        self.assertTrue(os.path.exists(src))
        self.assertTrue(os.path.exists(dst))

    def test_delete_file(self):
        from src.tools.file_tools import write_file, delete_file
        path = os.path.join(self.tmp, "todel.txt")
        write_file(path, "delete me")
        result = delete_file(path)
        self.assertTrue(result["success"])
        self.assertFalse(os.path.exists(path))

    def test_delete_nonexistent(self):
        from src.tools.file_tools import delete_file
        result = delete_file("/no/such/file.txt")
        self.assertIn("error", result)

    def test_get_file_info(self):
        from src.tools.file_tools import write_file, get_file_info
        path = os.path.join(self.tmp, "info.txt")
        write_file(path, "hello world")
        result = get_file_info(path)
        self.assertNotIn("error", result)
        self.assertEqual(result["type"], "file")
        self.assertEqual(result["extension"], ".txt")
        self.assertIn("md5", result)

    def test_create_directory(self):
        from src.tools.file_tools import create_directory
        path = os.path.join(self.tmp, "new", "nested", "dir")
        result = create_directory(path)
        self.assertTrue(result["success"])
        self.assertTrue(os.path.isdir(path))

    def test_tool_execute_interface(self):
        from src.tools import READ_FILE_TOOL, WRITE_FILE_TOOL
        path = os.path.join(self.tmp, "exec_test.txt")

        write_result = json.loads(WRITE_FILE_TOOL.execute(json.dumps({
            "path": path, "content": "exec interface test"
        })))
        self.assertTrue(write_result["success"])

        read_result = json.loads(READ_FILE_TOOL.execute(json.dumps({"path": path})))
        self.assertEqual(read_result["content"], "exec interface test")


class TestEditFileTools(unittest.TestCase):
    """Integration tests for edit_file tool in file_tools.py."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        # Reset the edit guard before each test to ensure isolation
        from src.core.edit_manager import reset_edit_guard
        reset_edit_guard()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)
        from src.core.edit_manager import reset_edit_guard
        reset_edit_guard()

    def test_edit_file_basic(self):
        """Full read → edit flow: read file, then apply a single SEARCH/REPLACE."""
        from src.tools.file_tools import write_file, read_file, edit_file
        path = os.path.join(self.tmp, "basic.py")
        write_file(path, "def foo():\n    return 1\n")

        # Must read first
        read_file(path)

        diff = (
            "------- SEARCH\n"
            "    return 1\n"
            "=======\n"
            "    return 42\n"
            "+++++++ REPLACE"
        )
        result = edit_file(path, diff)
        self.assertTrue(result["success"])
        self.assertEqual(result["applied_edits"], 1)

        # Verify file content
        with open(path, "r") as f:
            content = f.read()
        self.assertIn("return 42", content)
        self.assertNotIn("return 1", content)

    def test_edit_file_without_read_error(self):
        """Editing without prior read_file should be denied."""
        from src.tools.file_tools import write_file, edit_file
        path = os.path.join(self.tmp, "no_read.py")
        write_file(path, "original content")

        diff = (
            "------- SEARCH\n"
            "original content\n"
            "=======\n"
            "new content\n"
            "+++++++ REPLACE"
        )
        result = edit_file(path, diff)
        self.assertIn("error", result)
        self.assertEqual(result["error_type"], "edit_denied")

        # File should be unchanged
        with open(path, "r") as f:
            self.assertEqual(f.read(), "original content")

    def test_edit_file_multi_block(self):
        """Multiple SEARCH/REPLACE blocks in a single call."""
        from src.tools.file_tools import write_file, read_file, edit_file
        path = os.path.join(self.tmp, "multi.py")
        write_file(path, "alpha\nbeta\ngamma\ndelta\n")
        read_file(path)

        diff = (
            "------- SEARCH\n"
            "alpha\n"
            "=======\n"
            "ALPHA\n"
            "+++++++ REPLACE\n"
            "------- SEARCH\n"
            "gamma\n"
            "=======\n"
            "GAMMA\n"
            "+++++++ REPLACE"
        )
        result = edit_file(path, diff)
        self.assertTrue(result["success"])
        self.assertEqual(result["applied_edits"], 2)

        with open(path, "r") as f:
            content = f.read()
        self.assertEqual(content, "ALPHA\nbeta\nGAMMA\ndelta\n")

    def test_edit_file_no_match_error(self):
        """SEARCH block that doesn't match file content → error with suggestions."""
        from src.tools.file_tools import write_file, read_file, edit_file
        path = os.path.join(self.tmp, "nomatch.py")
        write_file(path, "def hello():\n    pass\n")
        read_file(path)

        diff = (
            "------- SEARCH\n"
            "def nonexistent():\n"
            "    pass\n"
            "=======\n"
            "def replaced():\n"
            "    pass\n"
            "+++++++ REPLACE"
        )
        result = edit_file(path, diff)
        self.assertIn("error", result)
        self.assertEqual(result["error_type"], "apply_failed")
        self.assertIn("failed_edits", result)

        # File should be unchanged
        with open(path, "r") as f:
            self.assertEqual(f.read(), "def hello():\n    pass\n")

    def test_edit_file_atomic_write(self):
        """After successful edit, file should contain the new content atomically."""
        from src.tools.file_tools import write_file, read_file, edit_file
        path = os.path.join(self.tmp, "atomic.txt")
        original = "line1\nline2\nline3\n"
        write_file(path, original)
        read_file(path)

        diff = (
            "------- SEARCH\n"
            "line2\n"
            "=======\n"
            "LINE_TWO\n"
            "+++++++ REPLACE"
        )
        result = edit_file(path, diff)
        self.assertTrue(result["success"])

        with open(path, "r") as f:
            content = f.read()
        self.assertEqual(content, "line1\nLINE_TWO\nline3\n")

        # No temp files should remain
        tmp_files = [f for f in os.listdir(self.tmp) if f.endswith(".tmp")]
        self.assertEqual(len(tmp_files), 0)

    def test_edit_file_tool_execute_interface(self):
        """Test edit_file via EDIT_FILE_TOOL.execute() JSON interface."""
        from src.tools.file_tools import write_file, read_file
        from src.tools import EDIT_FILE_TOOL
        path = os.path.join(self.tmp, "execute_test.txt")
        write_file(path, "hello world")
        read_file(path)

        diff = (
            "------- SEARCH\n"
            "hello world\n"
            "=======\n"
            "hello universe\n"
            "+++++++ REPLACE"
        )
        result = json.loads(EDIT_FILE_TOOL.execute(json.dumps({
            "file_path": path,
            "diff": diff,
        })))
        self.assertTrue(result["success"])

        with open(path, "r") as f:
            self.assertEqual(f.read(), "hello universe")


class TestBashTools(unittest.TestCase):
    """Tests for bash_tools.py."""

    def test_run_command_echo(self):
        from src.tools.bash_tools import run_command
        result = run_command("echo hello_agent_x1")
        self.assertTrue(result["success"])
        self.assertIn("hello_agent_x1", result["stdout"])
        self.assertEqual(result["return_code"], 0)

    def test_run_command_exit_code(self):
        from src.tools.bash_tools import run_command
        result = run_command("exit 42", timeout=5)
        self.assertFalse(result["success"])
        self.assertEqual(result["return_code"], 42)

    def test_run_command_blocked(self):
        from src.tools.bash_tools import run_command
        result = run_command("rm -rf /")
        self.assertIn("error", result)
        self.assertIn("blocked", result["error"].lower())

    def test_run_command_with_env(self):
        from src.tools.bash_tools import run_command
        result = run_command("echo $MY_TEST_VAR", env_vars={"MY_TEST_VAR": "agent_value"})
        self.assertTrue(result["success"])
        self.assertIn("agent_value", result["stdout"])

    def test_run_python_script(self):
        from src.tools.bash_tools import run_python_script
        with tempfile.NamedTemporaryFile(mode="w", suffix=".py", delete=False) as f:
            f.write("print('python_script_output')\n")
            script_path = f.name
        try:
            result = run_python_script(script_path)
            self.assertTrue(result["success"])
            self.assertIn("python_script_output", result["stdout"])
        finally:
            os.unlink(script_path)

    def test_run_python_script_not_found(self):
        from src.tools.bash_tools import run_python_script
        result = run_python_script("/no/such/script.py")
        self.assertIn("error", result)

    def test_run_bash_script(self):
        from src.tools.bash_tools import run_bash_script
        result = run_bash_script("echo bash_inline_test\necho line2")
        self.assertTrue(result["success"])
        self.assertIn("bash_inline_test", result["stdout"])
        self.assertIn("line2", result["stdout"])

    def test_get_system_info(self):
        from src.tools.bash_tools import get_system_info
        result = get_system_info()
        self.assertIn("os", result)
        self.assertIn("python_version", result)
        self.assertIn("cpu_count", result)
        self.assertIn("memory", result)

    def test_get_environment_variable(self):
        from src.tools.bash_tools import get_environment_variable
        os.environ["TEST_PLAIN_VAR"] = "testvalue"
        result = get_environment_variable("TEST_PLAIN_VAR")
        self.assertTrue(result["exists"])
        self.assertEqual(result["value"], "testvalue")

    def test_get_environment_variable_missing(self):
        from src.tools.bash_tools import get_environment_variable
        result = get_environment_variable("NONEXISTENT_VAR_XYZ_12345")
        self.assertFalse(result["exists"])
        self.assertIsNone(result["value"])

    def test_get_environment_variable_masked(self):
        from src.tools.bash_tools import get_environment_variable
        os.environ["MY_API_KEY"] = "super_secret"
        result = get_environment_variable("MY_API_KEY")
        self.assertTrue(result["exists"])
        self.assertEqual(result["value"], "***")

    def test_tool_execute_interface(self):
        from src.tools import RUN_COMMAND_TOOL
        result = json.loads(RUN_COMMAND_TOOL.execute(json.dumps({"command": "echo test_exec"})))
        self.assertTrue(result["success"])
        self.assertIn("test_exec", result["stdout"])


class TestPDFTools(unittest.TestCase):
    """Tests for pdf_tools.py."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _create_test_pdf(self, filename="test.pdf", text="Hello PDF World"):
        from src.tools.pdf_tools import create_pdf_from_text
        path = os.path.join(self.tmp, filename)
        result = create_pdf_from_text(text, path, title="Test PDF")
        return path, result

    def test_create_pdf_from_text(self):
        path, result = self._create_test_pdf()
        self.assertTrue(result["success"])
        self.assertTrue(os.path.exists(path))
        self.assertGreater(result["file_size_bytes"], 0)

    def test_read_pdf(self):
        path, _ = self._create_test_pdf(text="Readable content here")
        from src.tools.pdf_tools import read_pdf
        result = read_pdf(path)
        self.assertNotIn("error", result)
        self.assertEqual(result["total_pages"], 1)
        self.assertIn(1, result["pages"])

    def test_read_pdf_not_found(self):
        from src.tools.pdf_tools import read_pdf
        result = read_pdf("/no/such/file.pdf")
        self.assertIn("error", result)

    def test_get_pdf_metadata(self):
        path, _ = self._create_test_pdf()
        from src.tools.pdf_tools import get_pdf_metadata
        result = get_pdf_metadata(path)
        self.assertNotIn("error", result)
        self.assertEqual(result["total_pages"], 1)
        self.assertIn("file_size_bytes", result)

    def test_merge_pdfs(self):
        from src.tools.pdf_tools import merge_pdfs
        p1, _ = self._create_test_pdf("a.pdf", "Page one content")
        p2, _ = self._create_test_pdf("b.pdf", "Page two content")
        out = os.path.join(self.tmp, "merged.pdf")
        result = merge_pdfs([p1, p2], out)
        self.assertTrue(result["success"])
        self.assertEqual(result["total_pages"], 2)
        self.assertTrue(os.path.exists(out))

    def test_split_pdf(self):
        from src.tools.pdf_tools import split_pdf, merge_pdfs
        p1, _ = self._create_test_pdf("s1.pdf", "Page 1")
        p2, _ = self._create_test_pdf("s2.pdf", "Page 2")
        p3, _ = self._create_test_pdf("s3.pdf", "Page 3")
        merged = os.path.join(self.tmp, "to_split.pdf")
        merge_pdfs([p1, p2, p3], merged)

        out_dir = os.path.join(self.tmp, "split_output")
        result = split_pdf(merged, out_dir, pages_per_file=1)
        self.assertNotIn("error", result)
        self.assertEqual(result["files_created"], 3)

    def test_tool_execute_interface(self):
        from src.tools import CREATE_PDF_FROM_TEXT_TOOL, READ_PDF_TOOL
        path = os.path.join(self.tmp, "exec_pdf.pdf")
        create_result = json.loads(CREATE_PDF_FROM_TEXT_TOOL.execute(json.dumps({
            "text": "Tool execute test", "output_path": path
        })))
        self.assertTrue(create_result["success"])

        read_result = json.loads(READ_PDF_TOOL.execute(json.dumps({"path": path})))
        self.assertNotIn("error", read_result)


class TestPPTTools(unittest.TestCase):
    """Tests for ppt_tools.py."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _make_pptx(self, filename="test.pptx"):
        from src.tools.ppt_tools import create_presentation
        slides = [
            {"title": "Slide One", "content": "First slide body"},
            {"title": "Slide Two", "content": ["Bullet A", "Bullet B", "Bullet C"]},
        ]
        path = os.path.join(self.tmp, filename)
        result = create_presentation(slides, path, title="Unit Test Pres")
        return path, result

    def test_create_presentation(self):
        path, result = self._make_pptx()
        self.assertTrue(result["success"])
        self.assertEqual(result["slide_count"], 2)
        self.assertTrue(os.path.exists(path))
        self.assertGreater(result["file_size_bytes"], 0)

    def test_create_presentation_themes(self):
        from src.tools.ppt_tools import create_presentation
        for theme in ["default", "dark", "blue"]:
            path = os.path.join(self.tmp, f"theme_{theme}.pptx")
            result = create_presentation([{"title": theme, "content": "body"}], path, theme=theme)
            self.assertTrue(result["success"], f"Theme {theme} failed")

    def test_read_presentation(self):
        path, _ = self._make_pptx()
        from src.tools.ppt_tools import read_presentation
        result = read_presentation(path)
        self.assertNotIn("error", result)
        self.assertEqual(result["slide_count"], 2)
        self.assertIn("slides", result)

    def test_read_presentation_not_found(self):
        from src.tools.ppt_tools import read_presentation
        result = read_presentation("/no/such/file.pptx")
        self.assertIn("error", result)

    def test_add_slide(self):
        path, _ = self._make_pptx()
        from src.tools.ppt_tools import add_slide, read_presentation
        result = add_slide(path, "New Slide", "New slide content")
        self.assertTrue(result["success"])
        self.assertEqual(result["total_slides"], 3)
        read_result = read_presentation(path)
        self.assertEqual(read_result["slide_count"], 3)

    def test_tool_execute_interface(self):
        from src.tools import CREATE_PRESENTATION_TOOL, READ_PRESENTATION_TOOL
        path = os.path.join(self.tmp, "exec_ppt.pptx")
        create_result = json.loads(CREATE_PRESENTATION_TOOL.execute(json.dumps({
            "slides": [{"title": "Test Slide", "content": "Test content"}],
            "output_path": path
        })))
        self.assertTrue(create_result["success"])

        read_result = json.loads(READ_PRESENTATION_TOOL.execute(json.dumps({"path": path})))
        self.assertNotIn("error", read_result)
        self.assertEqual(read_result["slide_count"], 1)


class TestDataTools(unittest.TestCase):
    """Tests for data_tools.py."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _make_csv(self, filename="data.csv"):
        path = os.path.join(self.tmp, filename)
        with open(path, "w") as f:
            f.write("name,age,score\n")
            f.write("Alice,30,95.5\n")
            f.write("Bob,25,82.0\n")
            f.write("Charlie,35,78.3\n")
        return path

    def test_read_csv(self):
        from src.tools.data_tools import read_csv
        path = self._make_csv()
        result = read_csv(path)
        self.assertNotIn("error", result)
        self.assertEqual(result["total_rows"], 3)
        self.assertIn("name", result["columns"])
        self.assertEqual(len(result["data"]), 3)

    def test_read_csv_not_found(self):
        from src.tools.data_tools import read_csv
        result = read_csv("/no/such/file.csv")
        self.assertIn("error", result)

    def test_analyze_dataframe(self):
        from src.tools.data_tools import analyze_dataframe
        path = self._make_csv()
        result = analyze_dataframe(path)
        self.assertNotIn("error", result)
        self.assertEqual(result["total_rows"], 3)
        self.assertIn("age", result["statistics"])
        self.assertEqual(result["statistics"]["age"]["type"], "numeric")
        self.assertIn("name", result["statistics"])
        self.assertEqual(result["statistics"]["name"]["type"], "categorical")

    def test_filter_csv_equals(self):
        from src.tools.data_tools import filter_csv
        path = self._make_csv()
        result = filter_csv(path, [{"column": "name", "operator": "==", "value": "Alice"}])
        self.assertNotIn("error", result)
        self.assertEqual(result["filtered_rows"], 1)
        self.assertEqual(result["data"][0]["name"], "Alice")

    def test_filter_csv_gt(self):
        from src.tools.data_tools import filter_csv
        path = self._make_csv()
        result = filter_csv(path, [{"column": "age", "operator": ">", "value": 28}])
        self.assertEqual(result["filtered_rows"], 2)

    def test_filter_csv_contains(self):
        from src.tools.data_tools import filter_csv
        path = self._make_csv()
        result = filter_csv(path, [{"column": "name", "operator": "contains", "value": "li"}])
        self.assertEqual(result["filtered_rows"], 2)

    def test_filter_csv_invalid_column(self):
        from src.tools.data_tools import filter_csv
        path = self._make_csv()
        result = filter_csv(path, [{"column": "nonexistent", "operator": "==", "value": "x"}])
        self.assertIn("error", result)

    def test_save_as_csv(self):
        from src.tools.data_tools import save_as_csv, read_csv
        data = [{"x": 1, "y": "a"}, {"x": 2, "y": "b"}, {"x": 3, "y": "c"}]
        out = os.path.join(self.tmp, "saved.csv")
        result = save_as_csv(data, out)
        self.assertTrue(result["success"])
        self.assertEqual(result["rows"], 3)
        read_result = read_csv(out)
        self.assertEqual(len(read_result["data"]), 3)

    def test_read_json_file(self):
        from src.tools.data_tools import read_json_file
        path = os.path.join(self.tmp, "test.json")
        data = [{"key": "value1"}, {"key": "value2"}]
        with open(path, "w") as f:
            json.dump(data, f)
        result = read_json_file(path)
        self.assertNotIn("error", result)
        self.assertEqual(result["type"], "array")
        self.assertEqual(result["length"], 2)

    def test_read_json_file_object(self):
        from src.tools.data_tools import read_json_file
        path = os.path.join(self.tmp, "obj.json")
        with open(path, "w") as f:
            json.dump({"foo": "bar", "count": 42}, f)
        result = read_json_file(path)
        self.assertEqual(result["type"], "object")
        self.assertIn("foo", result["keys"])

    def test_convert_csv_to_json(self):
        from src.tools.data_tools import convert_data_format
        csv_path = self._make_csv()
        json_out = os.path.join(self.tmp, "converted.json")
        result = convert_data_format(csv_path, json_out, "csv", "json")
        self.assertTrue(result["success"])
        self.assertTrue(os.path.exists(json_out))

    def test_read_excel(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not available")
        import pandas as pd
        path = os.path.join(self.tmp, "test.xlsx")
        df = pd.DataFrame({"col1": [1, 2, 3], "col2": ["a", "b", "c"]})
        df.to_excel(path, index=False)
        from src.tools.data_tools import read_excel
        result = read_excel(path)
        self.assertNotIn("error", result)
        self.assertEqual(len(result["data"]), 3)

    def test_tool_execute_interface(self):
        from src.tools import READ_CSV_TOOL, ANALYZE_DATAFRAME_TOOL
        path = self._make_csv()
        read_result = json.loads(READ_CSV_TOOL.execute(json.dumps({"path": path})))
        self.assertNotIn("error", read_result)
        self.assertEqual(read_result["total_rows"], 3)

        analyze_result = json.loads(ANALYZE_DATAFRAME_TOOL.execute(json.dumps({"path": path})))
        self.assertNotIn("error", analyze_result)


class TestEconomicsTools(unittest.TestCase):
    """Tests for economics_tools.py - tests that don't require API keys."""

    def test_get_economic_calendar_no_key(self):
        from src.tools.economics_tools import get_economic_calendar
        os.environ.pop("TRADING_ECONOMICS_API_KEY", None)
        result = get_economic_calendar()
        self.assertNotIn("error", result)
        self.assertIn("note", result)
        self.assertIn("common_us_releases", result)
        self.assertGreater(len(result["common_us_releases"]), 0)

    def test_get_fred_no_key(self):
        from src.tools.economics_tools import get_fred_series
        os.environ.pop("FRED_API_KEY", None)
        result = get_fred_series("GDP")
        self.assertIn("error", result)
        self.assertIn("FRED_API_KEY", result["error"])

    def test_tool_schemas_valid(self):
        from src.tools import (
            GET_FRED_SERIES_TOOL, GET_WORLD_BANK_INDICATOR_TOOL,
            GET_EXCHANGE_RATES_TOOL, GET_ECONOMIC_CALENDAR_TOOL,
            GENERATE_ECONOMIC_REPORT_TOOL
        )
        for tool in [GET_FRED_SERIES_TOOL, GET_WORLD_BANK_INDICATOR_TOOL,
                     GET_EXCHANGE_RATES_TOOL, GET_ECONOMIC_CALENDAR_TOOL,
                     GENERATE_ECONOMIC_REPORT_TOOL]:
            schema = tool.get_schema()
            self.assertIn("function", schema)
            self.assertIsInstance(schema["function"]["description"], str)
            self.assertGreater(len(schema["function"]["description"]), 10)


class TestWebTools(unittest.TestCase):
    """Tests for web_tools.py - tests that don't require network."""

    def test_tool_schemas_valid(self):
        from src.tools import (
            FETCH_URL_TOOL, EXTRACT_WEBPAGE_TEXT_TOOL, EXTRACT_LINKS_TOOL,
            DOWNLOAD_FILE_TOOL, CHECK_URL_TOOL, FETCH_RSS_FEED_TOOL
        )
        for tool in [FETCH_URL_TOOL, EXTRACT_WEBPAGE_TEXT_TOOL, EXTRACT_LINKS_TOOL,
                     DOWNLOAD_FILE_TOOL, CHECK_URL_TOOL, FETCH_RSS_FEED_TOOL]:
            schema = tool.get_schema()
            self.assertIn("function", schema)

    def test_tool_execute_bad_url(self):
        from src.tools import CHECK_URL_TOOL
        result = json.loads(CHECK_URL_TOOL.execute(json.dumps({"url": "http://localhost:19999"})))
        self.assertFalse(result["reachable"])
        self.assertIn("error", result)

    def test_fetch_url_execute_invalid(self):
        from src.tools import FETCH_URL_TOOL
        result = json.loads(FETCH_URL_TOOL.execute(json.dumps({"url": "http://localhost:19999"})))
        self.assertIn("error", result)


class TestExistingToolsBackwardsCompatibility(unittest.TestCase):
    """Ensure existing tools still work after refactor."""

    def test_calculator_tool(self):
        from src.tools import CALCULATOR_TOOL
        result = json.loads(CALCULATOR_TOOL.execute('{"expression": "3 * 7"}'))
        self.assertEqual(result["result"], 21)

    def test_time_tool(self):
        from src.tools import TIME_TOOL
        result = json.loads(TIME_TOOL.execute('{}'))
        self.assertIn("datetime", result)
        self.assertIn("weekday", result)

    def test_weather_tool(self):
        from src.tools import WEATHER_TOOL
        result = json.loads(WEATHER_TOOL.execute('{"location": "Beijing"}'))
        self.assertIn("temperature", result)
        self.assertIn("condition", result)

    def test_all_tools_count(self):
        from src.tools import ALL_TOOLS
        self.assertGreaterEqual(len(ALL_TOOLS), 40)


if __name__ == "__main__":
    loader = unittest.TestLoader()
    suite = unittest.TestSuite()

    suite.addTests(loader.loadTestsFromTestCase(TestToolRegistry))
    suite.addTests(loader.loadTestsFromTestCase(TestFileTools))
    suite.addTests(loader.loadTestsFromTestCase(TestBashTools))
    suite.addTests(loader.loadTestsFromTestCase(TestPDFTools))
    suite.addTests(loader.loadTestsFromTestCase(TestPPTTools))
    suite.addTests(loader.loadTestsFromTestCase(TestDataTools))
    suite.addTests(loader.loadTestsFromTestCase(TestEconomicsTools))
    suite.addTests(loader.loadTestsFromTestCase(TestWebTools))
    suite.addTests(loader.loadTestsFromTestCase(TestExistingToolsBackwardsCompatibility))

    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)
    sys.exit(0 if result.wasSuccessful() else 1)
